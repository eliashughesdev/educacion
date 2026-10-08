import csv
import io
import json
from datetime import datetime, timezone
from html import escape

from fastapi import HTTPException
from fastapi.responses import Response

FORMATS = {
    "pdf": ("application/pdf", "pdf"),
    "xlsx": ("application/vnd.openxmlformats-officedocument.spreadsheetml.sheet", "xlsx"),
    "csv": ("text/csv; charset=utf-8", "csv"),
    "json": ("application/json", "json"),
    "docx": ("application/vnd.openxmlformats-officedocument.wordprocessingml.document", "docx"),
    "html": ("text/html; charset=utf-8", "html"),
}


def _txt(v) -> str:
    if v is None:
        return ""
    if isinstance(v, bool):
        return "Sí" if v else "No"
    if isinstance(v, float):
        return f"{v:,.2f}"
    if isinstance(v, (list, dict)):
        return json.dumps(v, ensure_ascii=False)
    return str(v)


def _stamp() -> str:
    return datetime.now(timezone.utc).strftime("%Y-%m-%d %H:%M UTC")


def _pdf(title, columns, rows, meta, totals) -> bytes:
    from reportlab.lib import colors
    from reportlab.lib.pagesizes import A4, landscape
    from reportlab.lib.styles import getSampleStyleSheet, ParagraphStyle
    from reportlab.platypus import SimpleDocTemplate, Table, TableStyle, Paragraph, Spacer

    buf = io.BytesIO()
    size = landscape(A4) if len(columns) > 5 else A4
    doc = SimpleDocTemplate(buf, pagesize=size, leftMargin=24, rightMargin=24, topMargin=28, bottomMargin=28, title=title)
    styles = getSampleStyleSheet()
    cell = ParagraphStyle("cell", parent=styles["BodyText"], fontSize=7.5, leading=9)
    head = ParagraphStyle("head", parent=cell, textColor=colors.white, fontName="Helvetica-Bold")
    story = [Paragraph(escape(title), styles["Title"])]
    for k, v in meta:
        story.append(Paragraph(f"<b>{escape(k)}:</b> {escape(_txt(v))}", styles["Normal"]))
    story.append(Spacer(1, 10))
    data = [[Paragraph(escape(c[1]), head) for c in columns]]
    data += [[Paragraph(escape(_txt(r.get(c[0]))), cell) for c in columns] for r in rows]
    if not rows:
        data.append([Paragraph("Sin registros", cell)] + [""] * (len(columns) - 1))
    table = Table(data, repeatRows=1)
    table.setStyle(TableStyle([
        ("BACKGROUND", (0, 0), (-1, 0), colors.HexColor("#0B1B33")),
        ("ROWBACKGROUNDS", (0, 1), (-1, -1), [colors.white, colors.HexColor("#EEF6FB")]),
        ("GRID", (0, 0), (-1, -1), 0.25, colors.HexColor("#9DB4C8")),
        ("VALIGN", (0, 0), (-1, -1), "TOP"),
    ]))
    story.append(table)
    for k, v in totals:
        story.append(Paragraph(f"<b>{escape(k)}:</b> {escape(_txt(v))}", styles["Normal"]))
    doc.build(story)
    return buf.getvalue()


def _xlsx(title, columns, rows, meta, totals) -> bytes:
    from openpyxl import Workbook
    from openpyxl.styles import Font, PatternFill

    wb = Workbook()
    ws = wb.active
    ws.title = title[:30].replace("/", "-")
    ws.append([title])
    ws["A1"].font = Font(bold=True, size=14)
    for k, v in meta:
        ws.append([k, _txt(v)])
    ws.append([])
    ws.append([c[1] for c in columns])
    hdr = ws.max_row
    for cell in ws[hdr]:
        cell.font = Font(bold=True, color="FFFFFF")
        cell.fill = PatternFill("solid", fgColor="0B1B33")
    for r in rows:
        ws.append([r.get(c[0]) if isinstance(r.get(c[0]), (int, float)) and not isinstance(r.get(c[0]), bool) else _txt(r.get(c[0])) for c in columns])
    for k, v in totals:
        ws.append([k, v])
    for i, c in enumerate(columns, start=1):
        ws.column_dimensions[ws.cell(row=hdr, column=i).column_letter].width = max(12, min(40, len(c[1]) + 6))
    ws.freeze_panes = ws.cell(row=hdr + 1, column=1)
    buf = io.BytesIO()
    wb.save(buf)
    return buf.getvalue()


def _csv(title, columns, rows, meta, totals) -> bytes:
    buf = io.StringIO()
    w = csv.writer(buf)
    w.writerow([c[1] for c in columns])
    for r in rows:
        w.writerow([r.get(c[0]) if isinstance(r.get(c[0]), (int, float)) else _txt(r.get(c[0])) for c in columns])
    return buf.getvalue().encode("utf-8-sig")


def _json(title, columns, rows, meta, totals) -> bytes:
    payload = {"title": title, "generated_at": _stamp(), "meta": dict(meta), "columns": [{"key": k, "label": l} for k, l in columns],
               "rows": [{k: r.get(k) for k, _ in columns} for r in rows], "totals": dict(totals), "count": len(rows)}
    return json.dumps(payload, ensure_ascii=False, indent=2, default=str).encode("utf-8")


def _docx(title, columns, rows, meta, totals) -> bytes:
    from docx import Document
    from docx.enum.section import WD_ORIENT
    from docx.shared import Pt

    d = Document()
    if len(columns) > 5:
        s = d.sections[0]
        s.orientation, s.page_width, s.page_height = WD_ORIENT.LANDSCAPE, s.page_height, s.page_width
    d.add_heading(title, level=1)
    for k, v in meta:
        d.add_paragraph(f"{k}: {_txt(v)}")
    t = d.add_table(rows=1, cols=len(columns))
    t.style = "Light Grid Accent 1"
    for i, c in enumerate(columns):
        t.rows[0].cells[i].text = c[1]
    for r in rows:
        cells = t.add_row().cells
        for i, c in enumerate(columns):
            cells[i].text = _txt(r.get(c[0]))
    for row in t.rows:
        for c in row.cells:
            for p in c.paragraphs:
                for run in p.runs:
                    run.font.size = Pt(8)
    for k, v in totals:
        d.add_paragraph(f"{k}: {_txt(v)}")
    buf = io.BytesIO()
    d.save(buf)
    return buf.getvalue()


def _html(title, columns, rows, meta, totals) -> bytes:
    head = "".join(f"<th>{escape(c[1])}</th>" for c in columns)
    body = "".join("<tr>" + "".join(f"<td>{escape(_txt(r.get(c[0])))}</td>" for c in columns) + "</tr>" for r in rows)
    metas = "".join(f"<div><b>{escape(k)}:</b> {escape(_txt(v))}</div>" for k, v in meta)
    tots = "".join(f"<div><b>{escape(k)}:</b> {escape(_txt(v))}</div>" for k, v in totals)
    html = f"""<!doctype html><html lang="es"><head><meta charset="utf-8"><title>{escape(title)}</title>
<style>body{{font-family:Arial,sans-serif;font-size:11px;margin:24px;color:#0b1b33}}h1{{font-size:18px}}table{{border-collapse:collapse;width:100%;margin-top:12px}}
th{{background:#0b1b33;color:#fff;text-align:left}}th,td{{border:1px solid #9db4c8;padding:4px 6px}}tr:nth-child(even) td{{background:#eef6fb}}
@media print{{@page{{size:{'landscape' if len(columns) > 5 else 'portrait'};margin:12mm}}}}</style></head>
<body onload="window.print()"><h1>{escape(title)}</h1>{metas}<table><thead><tr>{head}</tr></thead><tbody>{body}</tbody></table>{tots}</body></html>"""
    return html.encode("utf-8")


BUILDERS = {"pdf": _pdf, "xlsx": _xlsx, "csv": _csv, "json": _json, "docx": _docx, "html": _html}


def export_response(fmt: str, filename: str, title: str, columns, rows, meta=(), totals=()) -> Response:
    if fmt not in BUILDERS:
        raise HTTPException(400, f"Formato no soportado. Use: {', '.join(BUILDERS)}")
    meta = list(meta) + [("Generado", _stamp())]
    content = BUILDERS[fmt](title, columns, rows, meta, list(totals))
    media, ext = FORMATS[fmt]
    disposition = "inline" if fmt == "html" else "attachment"
    return Response(content, media_type=media, headers={"Content-Disposition": f'{disposition}; filename="{filename}.{ext}"'})
